"""
Independent verification of the Excel workbook's formula-driven sheets.

This does NOT just check that openpyxl wrote valid formula strings — it:
  1. Regenerates the source data independently.
  2. Builds the workbook.
  3. Forces LibreOffice (a real spreadsheet engine, not openpyxl) to open
     and recalculate every formula, headless.
  4. Reads back the recalculated values and compares them against an
     independent pandas computation of the same aggregates.

If openpyxl had written a formula Excel/LibreOffice couldn't parse, or
the SUMIFS/SUMPRODUCT logic didn't actually match the intended
aggregation, this test would fail on the recalculated numbers -- a
static "does the formula string look right" check would not catch that.
"""
import os
import subprocess
import sys
import tempfile

import openpyxl
import pandas as pd
import pytest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "src"))
from generate_data import generate_rows, COST_CENTERS, MONTHS  # noqa: E402
from build_workbook import build  # noqa: E402


@pytest.fixture(scope="module")
def recalculated_workbook():
    """Build the workbook, force a real recalc via LibreOffice headless,
    and return the path to the recalculated file plus the source
    DataFrame for independent comparison."""
    tmpdir = tempfile.mkdtemp()
    xlsx_path = os.path.join(tmpdir, "controlling_report_2025_2026.xlsx")
    build(output_path=xlsx_path, seed=42)

    recalced_dir = os.path.join(tmpdir, "recalced")
    os.makedirs(recalced_dir, exist_ok=True)
    result = subprocess.run(
        ["libreoffice", "--headless", "--convert-to",
         "xlsx:Calc MS Excel 2007 XML", "--outdir", recalced_dir, xlsx_path],
        capture_output=True, text=True, timeout=90,
    )
    recalced_path = os.path.join(recalced_dir, "controlling_report_2025_2026.xlsx")
    assert os.path.exists(recalced_path), (
        f"LibreOffice recalc failed.\nstdout: {result.stdout}\nstderr: {result.stderr}"
    )

    rows = generate_rows(seed=42)
    df = pd.DataFrame([r.__dict__ for r in rows])

    return recalced_path, df


def test_recalculated_file_opens_with_computed_values(recalculated_workbook):
    recalced_path, _ = recalculated_workbook
    wb = openpyxl.load_workbook(recalced_path, data_only=True)
    assert set(wb.sheetnames) == {
        "Rohdaten", "Kostenstellen-Uebersicht", "Monatsverlauf",
        "YoY-Vergleich", "Abweichungs-Report",
    }


def test_cost_center_totals_match_independent_pandas_computation(recalculated_workbook):
    recalced_path, df = recalculated_workbook
    wb = openpyxl.load_workbook(recalced_path, data_only=True)
    ws = wb["Kostenstellen-Uebersicht"]

    expected = df.groupby("cost_center_id")[["budget_eur", "actual_eur"]].sum()

    checked = 0
    for row in ws.iter_rows(min_row=2, max_row=1 + len(COST_CENTERS), values_only=True):
        cc_id, name, division, budget, actual, dev_eur, dev_pct = row
        exp_budget = expected.loc[cc_id, "budget_eur"]
        exp_actual = expected.loc[cc_id, "actual_eur"]
        assert budget == pytest.approx(exp_budget, abs=0.01), cc_id
        assert actual == pytest.approx(exp_actual, abs=0.01), cc_id
        assert dev_eur == pytest.approx(exp_actual - exp_budget, abs=0.01), cc_id
        checked += 1
    assert checked == len(COST_CENTERS)


def test_grand_total_row_matches_full_dataset_sum(recalculated_workbook):
    recalced_path, df = recalculated_workbook
    wb = openpyxl.load_workbook(recalced_path, data_only=True)
    ws = wb["Kostenstellen-Uebersicht"]

    total_row_idx = len(COST_CENTERS) + 2
    total_budget = ws.cell(row=total_row_idx, column=4).value
    total_actual = ws.cell(row=total_row_idx, column=5).value

    assert total_budget == pytest.approx(df["budget_eur"].sum(), abs=0.5)
    assert total_actual == pytest.approx(df["actual_eur"].sum(), abs=0.5)


def test_monthly_totals_match_independent_pandas_computation(recalculated_workbook):
    recalced_path, df = recalculated_workbook
    wb = openpyxl.load_workbook(recalced_path, data_only=True)
    ws = wb["Monatsverlauf"]

    expected = df.groupby("period")[["budget_eur", "actual_eur"]].sum()

    checked = 0
    for row in ws.iter_rows(min_row=2, max_row=1 + len(MONTHS), values_only=True):
        period, budget, actual, dev_eur, dev_pct = row
        assert budget == pytest.approx(expected.loc[period, "budget_eur"], abs=0.01), period
        assert actual == pytest.approx(expected.loc[period, "actual_eur"], abs=0.01), period
        checked += 1
    assert checked == len(MONTHS)


def test_yoy_sheet_correctly_isolates_2025_vs_2026(recalculated_workbook):
    recalced_path, df = recalculated_workbook
    wb = openpyxl.load_workbook(recalced_path, data_only=True)
    ws = wb["YoY-Vergleich"]

    df["year"] = df["period"].str[:4]
    expected = df.groupby(["cost_center_id", "year"])["actual_eur"].sum()

    checked = 0
    for row in ws.iter_rows(min_row=2, max_row=1 + len(COST_CENTERS), values_only=True):
        cc_id, name, actual_2025, actual_2026, change_eur, change_pct = row
        assert actual_2025 == pytest.approx(expected.loc[(cc_id, "2025")], abs=0.01), cc_id
        assert actual_2026 == pytest.approx(expected.loc[(cc_id, "2026")], abs=0.01), cc_id
        assert change_eur == pytest.approx(actual_2026 - actual_2025, abs=0.01), cc_id
        checked += 1
    assert checked == len(COST_CENTERS)


def test_marketing_cost_center_shows_genuine_underspend_signal(recalculated_workbook):
    """Cost center 6000 (Marketing) was deliberately generated with a
    consistent ~12% underspend factor -- this is the real signal the
    variance analysis exists to surface, not noise."""
    recalced_path, _ = recalculated_workbook
    wb = openpyxl.load_workbook(recalced_path, data_only=True)
    ws = wb["Kostenstellen-Uebersicht"]

    for row in ws.iter_rows(min_row=2, max_row=1 + len(COST_CENTERS), values_only=True):
        cc_id, name, division, budget, actual, dev_eur, dev_pct = row
        if cc_id == "6000":
            assert dev_pct < -0.08, f"Expected a genuine underspend signal, got {dev_pct}"
            return
    pytest.fail("Cost center 6000 not found")


def test_exception_report_only_contains_rows_above_threshold(recalculated_workbook):
    recalced_path, _ = recalculated_workbook
    wb = openpyxl.load_workbook(recalced_path, data_only=True)
    ws = wb["Abweichungs-Report"]

    rows = list(ws.iter_rows(min_row=2, values_only=True))
    rows = [r for r in rows if r[0] is not None]
    assert len(rows) > 0

    for period, cc_name, acc_name, budget, actual, dev_eur, dev_pct in rows:
        assert abs(dev_pct) > 0.15 - 1e-9
        assert abs(dev_eur) > 500 - 1e-6


def test_exception_report_is_sorted_by_absolute_deviation_descending(recalculated_workbook):
    recalced_path, _ = recalculated_workbook
    wb = openpyxl.load_workbook(recalced_path, data_only=True)
    ws = wb["Abweichungs-Report"]

    rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[0] is not None]
    deviations = [abs(r[5]) for r in rows]
    assert deviations == sorted(deviations, reverse=True)


def test_equipment_repair_spike_appears_in_exception_report(recalculated_workbook):
    """A one-off equipment repair spike was injected for cost center 1000
    ('Produktion Backwaren'), account 66000 ('Sonstige betriebliche
    Aufwendungen'), period 2026-03 -- this must surface as a flagged
    exception, proving the report catches real anomalies, not just
    aggregate drift."""
    recalced_path, _ = recalculated_workbook
    wb = openpyxl.load_workbook(recalced_path, data_only=True)
    ws = wb["Abweichungs-Report"]

    rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[0] is not None]
    match = [r for r in rows if r[0] == "2026-03" and r[1] == "Produktion Backwaren"
             and r[2] == "Sonstige betriebliche Aufwendungen"]
    assert len(match) == 1, "Expected the injected equipment-repair spike to be flagged"
    assert match[0][6] > 1.0, "Expected a large positive % deviation for the spike"


def test_raw_data_row_count_matches_generator(recalculated_workbook):
    recalced_path, df = recalculated_workbook
    wb = openpyxl.load_workbook(recalced_path, data_only=True)
    ws = wb["Rohdaten"]
    n_rows = sum(1 for row in ws.iter_rows(min_row=2, values_only=True) if row[0] is not None)
    assert n_rows == len(df)
